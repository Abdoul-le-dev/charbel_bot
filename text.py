"""
Bot Trading Pour Tous — webinaire des 30 septembre et 1er octobre 2026.
Fichier principal.   Lancer :  BOT_TOKEN=xxxx python script.py
"""
import asyncio
import functools
import html
import io
import os
import re
import traceback
import unicodedata
from collections import Counter
from datetime import datetime, timedelta, timezone
from urllib.parse import urlencode

from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    Application, CallbackQueryHandler, ChatJoinRequestHandler, CommandHandler,
    ContextTypes, ConversationHandler, MessageHandler, filters,
)

from database import database as db
from sondage import init_sondage_db, register_sondage_handlers   # module sondage : inchangé

# ════════════════════════════════════════════════════════════════════════════
# CONFIGURATION
# ════════════════════════════════════════════════════════════════════════════

TOKEN = os.getenv("BOT_TOKEN")          # ne JAMAIS écrire le token dans le code
ADMIN_IDS = {6992809421, 6799962131}
ADMIN_USERNAME = "@Faiseur2Rois"

WEBINAIRE = "2026-09-30"                # identifiant du webinaire : sert à savoir qui est déjà inscrit
HEURE_LIVE = 21                         # heure du Bénin
DUREE_LIVE = 2                          # heures (uniquement pour le calendrier)
JOURS = {
    1: {"date": "2026-09-30", "nom": "30 septembre", "semaine": "mercredi",
        "live": "https://youtube.com/live/jEC7hAcVlTQ?feature=share"},
    2: {"date": "2026-10-01", "nom": "1er octobre", "semaine": "jeudi",
        "live": "https://youtube.com/live/TNjFGH64FHQ?feature=share"},
}

VIDEO_BIENVENUE = "video/welcomes.MP4"
VIDEO_RELANCE = "bienvenu.mp4"                    # à la racine
# VIDEO_BULLE = "bulle.mp4"                       # vidéo ronde (voir demarrer_questionnaire)

FREINS = ["Le manque de temps", "La peur de perdre de l'argent",
          "Je ne sais pas par où commencer", "Autre"]

TOLERANCE = timedelta(minutes=20)       # au-delà, un rappel en retard n'est plus envoyé
PAUSE_ENVOI = 0.1                       # secondes entre deux envois (limite Telegram)
A_COMPLETER = "À COMPLÉTER"             # un rappel contenant ce mot n'est jamais envoyé

FORMAT = "%Y-%m-%d %H:%M:%S"
h = html.escape


def debut_live(jour) -> datetime:
    """Début du live, heure du Bénin."""
    return datetime.strptime(JOURS[jour]["date"], "%Y-%m-%d") + timedelta(hours=HEURE_LIVE)


SEUILS = {j: debut_live(j).strftime(FORMAT) for j in JOURS}    # clic à partir de là = présent


def jour_courant():
    aujourdhui = db.now_benin().strftime("%Y-%m-%d")
    return next((j for j, v in JOURS.items() if v["date"] == aujourdhui), None)


def dates_texte(jours) -> str:
    return "le " + " et le ".join(JOURS[j]["nom"] for j in jours)


def jours_choisis(u) -> list[int]:
    return [j for j in JOURS if u and u.get(f"veut_j{j}")] or list(JOURS)


# ════════════════════════════════════════════════════════════════════════════
# OUTILS D'ENVOI
# ════════════════════════════════════════════════════════════════════════════

TACHES = set()          # garde une référence aux tâches de fond


def lancer(coroutine):
    tache = asyncio.create_task(coroutine)
    TACHES.add(tache)
    tache.add_done_callback(TACHES.discard)


def kb(*boutons):
    """Un bouton par ligne. Chaque bouton = (texte, callback_data)."""
    return InlineKeyboardMarkup([[InlineKeyboardButton(t, callback_data=d)] for t, d in boutons])


def kb_demarrer(texte="Je confirme ma place"):
    return kb((texte, "demarrer"))


async def ecrire(bot, uid, texte, markup=None):
    await bot.send_message(chat_id=uid, text=texte, parse_mode="HTML", reply_markup=markup)


async def envoyer_video(bot, uid, chemin, caption=None, markup=None, nom=None, parse_mode=None):
    """Envoie une vidéo ; le fichier n'est téléversé qu'une fois (file_id gardé en base)."""
    nom = nom or chemin
    file_id = db.get_file_id(nom)
    if file_id:
        await bot.send_video(chat_id=uid, video=file_id, caption=caption,
                             parse_mode=parse_mode, reply_markup=markup)
        return
    with open(chemin, "rb") as f:
        msg = await bot.send_video(chat_id=uid, video=f, caption=caption,
                                   parse_mode=parse_mode, reply_markup=markup)
    db.save_file_id(nom, msg.video.file_id)


async def prevenir_admins(bot, texte):
    for admin_id in ADMIN_IDS:
        try:
            await bot.send_message(chat_id=admin_id, text=texte)
        except Exception as e:
            print(f"Notif admin {admin_id} : {e}")


def admin_seulement(fonction):
    @functools.wraps(fonction)
    async def enveloppe(update, ctx, *args, **kwargs):
        if update.effective_user.id not in ADMIN_IDS:
            await update.message.reply_text("⛔ Commande réservée à l'administrateur.")
            return ConversationHandler.END
        return await fonction(update, ctx, *args, **kwargs)
    return enveloppe


# ════════════════════════════════════════════════════════════════════════════
# NETTOYAGE DU PRÉNOM (logique d'origine, inchangée)
# ════════════════════════════════════════════════════════════════════════════

def _nettoyer_prenom(texte: str) -> str:
    t = texte.strip()
    t = t.replace("\u2019", "'").replace("\u2018", "'")
    t = t.lower()
    t = "".join(c for c in t if unicodedata.category(c) not in ("So", "Sm", "Sk", "Cn"))

    for ancien, nouveau in [
        ("mappelle", " "), ("cest", " "), ("mest", " "),
        ("jsuis", " "), ("chuis", " "),
    ]:
        t = t.replace(ancien, nouveau)

    parasites = [
        "tout le monde m'appelle", "vous pouvez m'appeler", "tu peux m'appeler",
        "ils m'appellent", "on m'appelle",
        "je me présente", "je me presente",
        "ravi de te rencontrer", "ravi de vous rencontrer",
        "comment ça va", "comment ca va",
        "bonne journée", "bonne soirée", "bonne nuit",
        "bien sûr", "bien sur",
        "je m'appelle", "j'me appelle",
        "mon prénom c'est", "mon prenom c'est",
        "mon prénom est", "mon prenom est",
        "mon nom c'est", "mon nom est",
        "je me nomme", "je me nome",
        "moi c'est", "c'est moi",
        "my name is", "my name's",
        "they call me", "call me",
        "people call me", "everyone calls me", "you can call me",
        "je suis", "je sui",
        "i am", "i'm",
        "avec plaisir",
        "bonjour", "bonsoir", "salut", "coucou", "hello",
        "hey", "wesh", "yo", "slt", "salu", "hi",
        "enchanté", "enchante", "enchantée",
        "voilà", "voila", "voici",
        "exactement", "exact",
        "ouais", "ouas", "oui",
        "alors", "donc", "bien",
        "prénom", "prenom", "appelle", "appeler", "nomme",
        "ça va", "ca va",
        "ok",
        "mon", "ma", "mes", "nom",
        "moi", "moa", "mwa",
        "suis", "est", "dj", "je", "et",
    ]

    for p in parasites:
        pattern = re.escape(p)
        t = re.sub(r"(?<![a-zà-ÿ])" + pattern + r"(?![a-zà-ÿ])", " ", t)

    t = re.sub(r"[^\w\s\u00C0-\u00FF\-']", " ", t)
    t = re.sub(r"(?<![a-zà-ÿ])'|'(?![a-zà-ÿ])", " ", t)
    t = re.sub(r"\s+", " ", t).strip()

    def cap_mot(m: str) -> str:
        return "-".join(p.capitalize() for p in m.split("-"))

    mots = [cap_mot(m) for m in t.split() if len(m) > 1]
    return " ".join(mots) if mots else texte.strip().split()[-1].capitalize()


def _extraire_prenom(texte: str) -> tuple[str, bool]:
    """Retourne (prénom, confirmation_requise)."""
    brut = texte.strip()
    if len(brut.split()) <= 2 and len(brut) <= 20:
        return brut.title(), False
    return _nettoyer_prenom(brut), True


def _valider_whatsapp(texte: str) -> str | None:
    """Accepte avec ou sans '+', espaces/tirets tolérés, 8 à 15 chiffres."""
    texte = texte.strip()
    if not re.fullmatch(r"\+?[\d\s.\-()]+", texte):
        return None
    chiffres = re.sub(r"\D", "", texte)
    if not 8 <= len(chiffres) <= 15:
        return None
    return ("+" if texte.startswith("+") else "") + chiffres


# ════════════════════════════════════════════════════════════════════════════
# QUESTIONNAIRE — piloté par la base : on repart toujours de la 1re question sans réponse
# ════════════════════════════════════════════════════════════════════════════

PRENOMS_A_CONFIRMER = {}        # uid -> prénom proposé, en attente du bouton Oui/Non


def prochaine_etape(u) -> str:
    if not u.get("prenom"):     return "prenom"
    if not u.get("whatsapp"):   return "whatsapp"
    if not u.get("deja_trade"): return "trade"
    if not u.get("frein"):      return "frein"
    if not u.get("presence"):   return "presence"
    if u["presence"] == "Un seul soir" and u.get("veut_j1") is None:
        return "soir"
    return "confirmation"


def restantes(u) -> int:
    manquantes = sum(1 for c in ("prenom", "whatsapp", "deja_trade", "frein", "presence") if not u.get(c))
    return manquantes + (prochaine_etape(u) == "soir")


def inscrit(u) -> bool:
    return bool(u and u.get("completed") == 1 and u.get("webinaire") == WEBINAIRE)


async def message_deja_inscrit(bot, uid):
    u = db.get_user(uid)
    await ecrire(bot, uid,
                 "<b>Tu es déjà inscrit(e)</b>\n\n"
                 f"Rendez-vous {dates_texte(jours_choisis(u))} à {HEURE_LIVE}h00 (heure du Bénin).\n\n"
                 "Ton lien du direct et les rappels arrivent ici, en privé.")


async def demarrer_questionnaire(bot, uid):
    u = db.get_user(uid)
    if inscrit(u):                                  # jamais deux fois pour la même personne
        await message_deja_inscrit(bot, uid)
        return

    if not u or u.get("webinaire") != WEBINAIRE:
        # Nouveau webinaire : on garde prénom / WhatsApp / Q3 / Q4 déjà connus,
        # mais la présence est à redonner (les anciennes réponses concernaient septembre).
        db.upsert_user(uid, webinaire=WEBINAIRE, presence=None, veut_j1=None, veut_j2=None,
                       completed=0, relance10=0, relance30=0)
    db.upsert_user(uid, en_cours=1)

    await ecrire(
        bot, uid,
        "<b>Confirme ta place</b>\nFormation gratuite Trading Pour Tous\n\n"
        f"{JOURS[1]['nom']} et {JOURS[2]['nom']}, {HEURE_LIVE}h00 (heure du Bénin)\n"
        "En direct uniquement, places limitées.")

    # ── Vidéo en bulle (désactivée) — pour l'activer : décommenter, mettre le fichier à la racine
    # (carré, 1 minute max) et définir VIDEO_BULLE en haut du fichier.
    # with open(VIDEO_BULLE, "rb") as f:
    #     await bot.send_video_note(chat_id=uid, video_note=f)

    await poser_question(bot, uid)


async def poser_question(bot, uid):
    u = db.get_user(uid)
    etape = prochaine_etape(u)

    if etape == "prenom":
        await ecrire(bot, uid, "<b>Question 1 sur 5</b>\n\nQuel est ton prénom ?")
    elif etape == "whatsapp":
        await ecrire(bot, uid, "<b>Question 2 sur 5</b>\n\nQuel est ton numéro WhatsApp ?\n\n"
                               "<i>Avec l'indicatif, par exemple : +229 60619292</i>")
    elif etape == "trade":
        await ecrire(bot, uid, "<b>Question 3 sur 5</b>\n\nAs-tu déjà fait du trading ?",
                     kb(("Oui", "q:trade:Oui"), ("Non", "q:trade:Non")))
    elif etape == "frein":
        await ecrire(bot, uid, "<b>Question 4 sur 5</b>\n\nQu'est-ce qui t'a empêché jusqu'ici de te lancer ?",
                     kb(*[(f, f"q:frein:{i}") for i, f in enumerate(FREINS)]))
    elif etape == "presence":
        await ecrire(bot, uid, "<b>Question 5 sur 5</b>\n\nTa présence\n\n<i>En direct uniquement, pas de replay.</i>",
                     kb((f"Je serai là en direct les DEUX soirs, à {HEURE_LIVE}h", "q:presence:deux"),
                        ("Je serai là un seul soir", "q:presence:un")))
    elif etape == "soir":
        await ecrire(bot, uid, "<b>Lequel des deux soirs ?</b>",
                     kb(*[(JOURS[j]["nom"], f"q:soir:{j}") for j in JOURS]))
    else:
        await ecrire(bot, uid, f"<b>Tout est prêt, {h(u.get('prenom') or '')}.</b>\n\n"
                               "Il ne reste qu'un clic pour verrouiller ta place.",
                     kb(("Confirmer ma place", "q:confirme:1")))


async def repondre_texte(bot, uid, texte) -> bool:
    """Texte reçu = réponse au questionnaire (prénom / WhatsApp) ? Retourne False si ce n'est pas le cas."""
    u = db.get_user(uid)
    if not u or not u.get("en_cours"):
        return False
    etape = prochaine_etape(u)

    if etape == "prenom":
        prenom, a_confirmer = _extraire_prenom(texte)
        if a_confirmer:
            if len(prenom) > 15:
                await ecrire(bot, uid, "J'ai du mal à identifier ton prénom dans ce message.\n\n"
                                       "Peux-tu m'envoyer <b>uniquement ton prénom</b> ?")
                return True
            PRENOMS_A_CONFIRMER[uid] = prenom
            await ecrire(bot, uid, f"Ton prénom est bien <b>{h(prenom)}</b> ?",
                         kb((f"Oui, c'est bien {prenom}", "prenom:oui"), ("Non, je corrige", "prenom:non")))
            return True
        db.upsert_user(uid, prenom=prenom)

    elif etape == "whatsapp":
        numero = _valider_whatsapp(texte)
        if not numero:
            await ecrire(bot, uid, "Ce numéro ne semble pas valide.\n\n"
                                   "Envoie-le avec l'indicatif, par exemple :\n<b>+229 60619292</b>")
            return True
        db.upsert_user(uid, whatsapp=numero)

    else:
        return False

    await poser_question(bot, uid)
    return True


async def confirmer_prenom(bot, uid, reponse) -> bool:
    u = db.get_user(uid)
    if not u or prochaine_etape(u) != "prenom":
        return False
    prenom = PRENOMS_A_CONFIRMER.pop(uid, None)
    if reponse == "oui" and prenom:
        db.upsert_user(uid, prenom=prenom)
        await poser_question(bot, uid)
    else:
        await ecrire(bot, uid, "Pas de souci.\n\nEnvoie-moi juste ton prénom :")
    return True


async def repondre_bouton(bot, uid, data) -> bool:
    """data = 'trade:Oui', 'frein:2', 'presence:deux', 'soir:1', 'confirme:1'.
    N'est accepté que si ça correspond à la question en cours (anti double-clic, anciens boutons)."""
    u = db.get_user(uid)
    champ, _, valeur = data.partition(":")
    if not u or inscrit(u) or prochaine_etape(u) != ("confirmation" if champ == "confirme" else champ):
        return False

    if champ == "trade":
        db.upsert_user(uid, deja_trade=valeur)
    elif champ == "frein":
        db.upsert_user(uid, frein=FREINS[int(valeur)])
    elif champ == "presence":
        if valeur == "deux":
            db.upsert_user(uid, presence="Les deux soirs", veut_j1=1, veut_j2=1)
        else:
            db.upsert_user(uid, presence="Un seul soir")
    elif champ == "soir":
        db.upsert_user(uid, veut_j1=int(valeur == "1"), veut_j2=int(valeur == "2"))
    elif champ == "confirme":
        db.upsert_user(uid, completed=1, en_cours=0)
        await envoyer_confirmation(bot, uid)
        return True

    await poser_question(bot, uid)
    return True


async def gerer_souci(bot, uid):
    """Bouton « Je suis bloqué(e) » : on donne le contact de l'admin et on met le questionnaire en pause
    (plus aucune relance automatique pour cette personne)."""
    db.upsert_user(uid, en_cours=0)
    await ecrire(bot, uid,
                 "<b>Pas de souci.</b>\n\n"
                 f"Pour t'aider rapidement, contacte directement Charbel : {ADMIN_USERNAME}\n\n"
                 "Quand tu es prêt(e), tu peux reprendre ton inscription.",
                 kb_demarrer("Reprendre mon inscription"))


# ════════════════════════════════════════════════════════════════════════════
# CONFIRMATION : place confirmée → calendrier → lien du live
# ════════════════════════════════════════════════════════════════════════════

def _utc(jour):
    """Début et fin du live en UTC (Bénin = UTC+1)."""
    debut = debut_live(jour) - timedelta(hours=1)
    return debut, debut + timedelta(hours=DUREE_LIVE)


def lien_google_agenda(jour) -> str:
    debut, fin = _utc(jour)
    return "https://calendar.google.com/calendar/render?" + urlencode({
        "action": "TEMPLATE",
        "text": "Formation gratuite Trading Pour Tous",
        "dates": f"{debut:%Y%m%dT%H%M%SZ}/{fin:%Y%m%dT%H%M%SZ}",
        "details": f"Formation en direct (heure du Bénin).\n\nLien du direct : {JOURS[jour]['live']}",
        "location": "En direct en ligne",
    }, safe="/")


def fichier_ics(jours) -> bytes:
    """Fichier calendrier pour iPhone : un événement par soir choisi, avec le lien du live dedans."""
    lignes = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Trading Pour Tous//FR", "CALSCALE:GREGORIAN"]
    for j in jours:
        debut, fin = _utc(j)
        lignes += [
            "BEGIN:VEVENT",
            f"UID:tpt-{WEBINAIRE}-j{j}@tradingpourtous",
            f"DTSTAMP:{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}",
            f"DTSTART:{debut:%Y%m%dT%H%M%SZ}",
            f"DTEND:{fin:%Y%m%dT%H%M%SZ}",
            "SUMMARY:Formation gratuite Trading Pour Tous",
            f"DESCRIPTION:Formation en direct (heure du Bénin).\\nLien du direct : {JOURS[j]['live']}",
            "LOCATION:En direct en ligne",
            f"URL:{JOURS[j]['live']}",
            "BEGIN:VALARM", "TRIGGER:-PT30M", "ACTION:DISPLAY",
            "DESCRIPTION:La formation commence dans 30 minutes", "END:VALARM",
            "END:VEVENT",
        ]
    lignes.append("END:VCALENDAR")
    return "\r\n".join(lignes).encode("utf-8")


async def envoyer_confirmation(bot, uid):
    u = db.get_user(uid)
    jours = jours_choisis(u)
    await ecrire(bot, uid, f"<b>C'est noté, {h(u['prenom'])}.</b>\n\n"
                           f"Ta place est confirmée pour {dates_texte(jours)} à {HEURE_LIVE}h.")
    await ecrire(bot, uid, "<b>Ajoute la formation à ton calendrier</b>\n\nPour ne pas oublier.",
                 kb(("Android : Google Agenda", "agenda:android"),
                    ("iPhone : fichier calendrier", "agenda:iphone")))
    await ecrire(bot, uid,
                 "<b>Ton lien pour le direct</b>\n\n"
                 "Garde-le précieusement.\n\n"
                 "Je te rappellerai ici même, en privé, la veille et le jour J.\n\n"
                 f"À {JOURS[jours[0]]['semaine']} soir !",
                 kb(*[(f"Live du {JOURS[j]['nom']}", f"live:{j}") for j in jours]))


async def envoyer_agenda(bot, uid, plateforme):
    """Clic sur un bouton calendrier : on note le clic, puis on envoie le lien (Android) ou le .ics (iPhone)."""
    db.enregistrer_suivi(uid, f"agenda_{plateforme}")
    jours = jours_choisis(db.get_user(uid))
    if plateforme == "android":
        bouton = InlineKeyboardMarkup([[InlineKeyboardButton("Ajouter à Google Agenda",
                                                              url=lien_google_agenda(jours[0]))]])
        await ecrire(bot, uid, "Clique ci-dessous pour l'ajouter à ton agenda.", bouton)
    else:
        await bot.send_document(chat_id=uid, document=fichier_ics(jours), filename="formation-trading-pour-tous.ics",
                                caption="Ouvre ce fichier pour l'ajouter à ton calendrier iPhone.")


async def envoyer_lien_live(bot, uid, jour):
    """Clic sur « Rejoindre le live » : on note le clic (jour + heure), puis on envoie le vrai lien."""
    db.enregistrer_suivi(uid, "live", jour)
    bouton = InlineKeyboardMarkup([[InlineKeyboardButton("Ouvrir le live", url=JOURS[jour]["live"])]])
    await ecrire(bot, uid, f"<b>Ton lien pour le direct du {JOURS[jour]['nom']}</b>", bouton)


# ════════════════════════════════════════════════════════════════════════════
# ARRIVÉE DANS LE CANAL ET /start
# ════════════════════════════════════════════════════════════════════════════

async def envoyer_bienvenue(bot, uid, lien_entree=None):
    db.log_member(uid, lien_entree)
    db.upsert_user(uid, categorie=db.derniere_categorie())

    legende = (
        "<b>Bienvenue dans Trading Pour Tous</b>\n\n"
        "Tu es sur le point de réserver ta place à la formation gratuite.\n\n"
        f"<b>Dates</b>\n{JOURS[1]['nom']} et {JOURS[2]['nom']}, à {HEURE_LIVE}h00 (heure du Bénin)\n\n"
        "<b>Format</b>\nEn direct uniquement. Places limitées.")
    await envoyer_video(bot, uid, VIDEO_BIENVENUE, legende, nom="welcomes_222", parse_mode="HTML")
    await ecrire(bot, uid, "<b>Il reste peu de places.</b>\n\nClique ci-dessous pour confirmer la tienne.",
                 kb_demarrer())


async def envoyer_bienvenue_securise(bot, uid, lien_entree=None):
    try:
        await envoyer_bienvenue(bot, uid, lien_entree)
    except Exception as e:
        print(f"Erreur bienvenue uid={uid} : {e}")


async def approuver_demande(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    demande = update.chat_join_request
    uid = demande.from_user.id
    try:
        await demande.approve()
    except Exception as e:
        print(f"approve() uid={uid} : {e}")
    lien = demande.invite_link          # nommer les liens A, B, C dans Telegram : le nom est gardé en base
    nom_lien = (lien.name or lien.invite_link) if lien else None
    lancer(envoyer_bienvenue_securise(ctx.bot, uid, nom_lien))


async def cmd_start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    db.touch_last_seen(uid)
    if inscrit(db.get_user(uid)):
        await message_deja_inscrit(ctx.bot, uid)
    else:
        lancer(envoyer_bienvenue_securise(ctx.bot, uid))


async def cmd_inscription(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    db.touch_last_seen(update.effective_user.id)
    await demarrer_questionnaire(ctx.bot, update.effective_user.id)


# ════════════════════════════════════════════════════════════════════════════
# BOUTONS (un seul routeur)
# ════════════════════════════════════════════════════════════════════════════

async def boutons(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    uid, bot = q.from_user.id, ctx.bot
    db.touch_last_seen(uid)
    action, _, reste = q.data.partition(":")

    accepte = None
    if action == "q":
        accepte = await repondre_bouton(bot, uid, reste)
    elif action == "prenom":
        accepte = await confirmer_prenom(bot, uid, reste)
    elif action == "live":
        await envoyer_lien_live(bot, uid, int(reste))
    elif action == "agenda":
        await envoyer_agenda(bot, uid, reste)
    elif action == "demarrer":
        await demarrer_questionnaire(bot, uid)
    elif action == "souci":
        await gerer_souci(bot, uid)

    if accepte:                          # on retire les boutons d'une question déjà répondue
        try:
            await q.edit_message_reply_markup(reply_markup=None)
        except Exception:
            pass


# ════════════════════════════════════════════════════════════════════════════
# MESSAGES LIBRES ET CODES DU SOIR
# ════════════════════════════════════════════════════════════════════════════

def normaliser(texte) -> str:
    """'  Nasdaq ! ' -> 'NASDAQ' (sans accents, sans ponctuation)."""
    texte = unicodedata.normalize("NFKD", texte)
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", texte.upper())


async def traiter_code(bot, uid, texte, jour_actuel=None) -> bool:
    """Le message est-il l'un des codes du soir ? Si oui on l'enregistre (qui + quelle heure)."""
    trouves = db.trouver_code(normaliser(texte))
    if not trouves:
        return False
    jour_actuel = jour_actuel or jour_courant()
    jour, rang = next(((j, r) for j, r in trouves if j == jour_actuel), trouves[0])
    db.enregistrer_suivi(uid, "code", jour, str(rang))
    prenom = (db.get_user(uid) or {}).get("prenom") or ""
    await ecrire(bot, uid, f"Code reçu, {h(prenom)}. Merci d'être en direct.")
    return True


async def message_libre(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if not update.message or not update.message.text:
        return
    user, bot = update.effective_user, ctx.bot
    uid, brut = user.id, update.message.text
    db.touch_last_seen(uid)

    if await traiter_code(bot, uid, brut):           # 1. code du soir
        return
    if await repondre_texte(bot, uid, brut):         # 2. réponse au questionnaire
        return

    db.log_message(uid, brut)                        # 3. message libre → on prévient l'admin
    texte = brut.lower().strip()
    if "present" in texte or "présent" in texte:
        u = db.get_user(uid)
        reponse = ("Je suis très content de savoir que tu seras là !\n\n"
                   f"N'oublie pas : {dates_texte(jours_choisis(u))} à {HEURE_LIVE}h, heure de Cotonou.\n\n"
                   "Je t'enverrai ton lien ici, en privé, juste avant le live.")
    elif "merci" in texte:
        reponse = "jtp !"
    elif "ok" in texte:
        reponse = "Super !"
    else:
        reponse = ("Ton message a bien été reçu.\n\n"
                   f"Pour une réponse rapide, contacte directement Charbel sur Telegram : {ADMIN_USERNAME}\n\n"
                   "Il te répondra dès que possible.")
    await update.message.reply_text(reponse)

    pseudo = f"@{user.username}" if user.username else f"id:{uid}"
    await prevenir_admins(bot, f"Nouveau message reçu\n\nUtilisateur : {user.first_name or ''} ({pseudo})\n"
                               f"ID : {uid}\n\nMessage : {brut[:200]}")


# ════════════════════════════════════════════════════════════════════════════
# CIBLES (utilisées par /envoyer, /export et les rappels)
# ════════════════════════════════════════════════════════════════════════════

def _incomplets():
    confirmes = set(db.ids_confirmes(WEBINAIRE))
    return [i for i in db.ids_tous() if i not in confirmes]


def _presents(jour):
    presents = set(db.ids_presents(jour, SEUILS[jour]))
    return [i for i in db.ids_confirmes(WEBINAIRE, jour) if i in presents]


def _absents(jour):
    presents = set(db.ids_presents(jour, SEUILS[jour]))
    return [i for i in db.ids_confirmes(WEBINAIRE, jour) if i not in presents]


# clé -> (libellé, fonction qui retourne la liste des identifiants)
CIBLES = {
    "tous":         ("Tous (canal + bot)",                    db.ids_tous),
    "confirmes":    ("Confirmés",                             lambda: db.ids_confirmes(WEBINAIRE)),
    "incomplets":   ("Incomplets (pas confirmés)",            _incomplets),
    "sept":         ("Anciens inscrits de septembre",         db.ids_septembre),
    "j1":           (f"Confirmés pour le {JOURS[1]['nom']}",  lambda: db.ids_confirmes(WEBINAIRE, 1)),
    "j2":           (f"Confirmés pour le {JOURS[2]['nom']}",  lambda: db.ids_confirmes(WEBINAIRE, 2)),
    "presents_1":   (f"Présents le {JOURS[1]['nom']}",        lambda: _presents(1)),
    "absents_1":    (f"Absents le {JOURS[1]['nom']}",         lambda: _absents(1)),
    "presents_2":   (f"Présents le {JOURS[2]['nom']}",        lambda: _presents(2)),
    "absents_2":    (f"Absents le {JOURS[2]['nom']}",         lambda: _absents(2)),
    "debutants":    ("Débutants (jamais tradé)",              lambda: db.ids_confirmes(WEBINAIRE, champ="deja_trade", valeur="Non")),
    "traders":      ("Ont déjà tradé",                        lambda: db.ids_confirmes(WEBINAIRE, champ="deja_trade", valeur="Oui")),
    "un_seul_soir": ("Un seul soir",                          lambda: db.ids_confirmes(WEBINAIRE, champ="presence", valeur="Un seul soir")),
}
for _i, _frein in enumerate(FREINS):
    CIBLES[f"frein_{_i}"] = (f"Frein : {_frein}",
                             lambda f=_frein: db.ids_confirmes(WEBINAIRE, champ="frein", valeur=f))


# ════════════════════════════════════════════════════════════════════════════
# RAPPELS PLANIFIÉS (modifiables par l'admin) + RELANCES AUTOMATIQUES
# ════════════════════════════════════════════════════════════════════════════

# (moment, texte) — moment = heure fixe "12:00" ou décalage en minutes par rapport au début du live.
# {prenom} est remplacé par le prénom. Le bouton « Rejoindre le live » est ajouté à chaque message.
SEQUENCE_SOIR = [
    ("12:00", f"{{prenom}}, ce soir à {HEURE_LIVE}h.\n\nTon lien est déjà là, juste en dessous."),
    (-180, "Dans 3 heures.\n\nCharge ton téléphone."),
    (-30, "Dans 30 minutes.\n\nConnecte-toi maintenant."),
    (5, "On a commencé.\n\nPlus d'une centaine de personnes sont déjà là."),
]


def _apres_live(jour, minutes) -> str:
    return (debut_live(jour) + timedelta(minutes=minutes)).strftime(FORMAT)


# (nom, cible, date d'envoi, texte, vidéo, bouton)  — bouton : "live_1" / "live_2" / "demarrer" / ""
PLANNING = [
    ("Réinvitation septembre", "sept", "2026-09-25 12:00:00",
     "{prenom},\n\nla formation gratuite revient :\nmercredi 30 septembre et jeudi 1er octobre, à "
     f"{HEURE_LIVE}h.\n\nTu avais réservé ta place en septembre. Elle t'attend toujours.\n\nConfirme-la ici.",
     "", "demarrer"),
    ("Veille J1 (vidéo)", "j1", "2026-09-29 20:00:00",
     f"Demain soir, {HEURE_LIVE}h.\n\nVoilà ce que tu vas voir.", "29.mp4", "live_1"),
    ("Fin J1 → demain", "confirmes", _apres_live(1, DUREE_LIVE * 60 + 15),
     f"Demain, je montre [{A_COMPLETER}].\n\nMême heure, même lien.", "", "live_2"),
    ("Absents J1 (replay)", "absents_1", _apres_live(1, DUREE_LIVE * 60 + 45),
     f"[{A_COMPLETER}] texte replay + relance", "", ""),
    ("Présents J1 (offre)", "presents_1", _apres_live(1, DUREE_LIVE * 60 + 45),
     f"[{A_COMPLETER}] texte de l'offre", "", ""),
    ("Absents J2 (replay)", "absents_2", _apres_live(2, DUREE_LIVE * 60 + 45),
     f"[{A_COMPLETER}] texte replay + relance", "", ""),
    ("Présents J2 (offre)", "presents_2", _apres_live(2, DUREE_LIVE * 60 + 45),
     f"[{A_COMPLETER}] texte de l'offre", "", ""),
]
for _jour, _infos in JOURS.items():
    for _moment, _texte in SEQUENCE_SOIR:
        if isinstance(_moment, str):
            _quand = datetime.strptime(f"{_infos['date']} {_moment}", "%Y-%m-%d %H:%M")
        else:
            _quand = debut_live(_jour) + timedelta(minutes=_moment)
        PLANNING.append((f"J{_jour} {_quand:%H:%M}", f"j{_jour}", _quand.strftime(FORMAT),
                         _texte, "", f"live_{_jour}"))
PLANNING.sort(key=lambda r: r[2])


def clavier_rappel(bouton):
    if bouton.startswith("live_"):
        return kb(("Rejoindre le live", f"live:{bouton[5:]}"))
    if bouton == "demarrer":
        return kb_demarrer()
    return None


async def envoyer_a(bot, uid, texte, video, markup):
    """Un message de rappel, en texte brut (le texte écrit par l'admin n'est jamais interprété)."""
    prenom = (db.get_user(uid) or {}).get("prenom") or "Hello"
    texte = texte.replace("{prenom}", prenom)
    if video:
        await envoyer_video(bot, uid, video, texte, markup)
    else:
        await bot.send_message(chat_id=uid, text=texte, reply_markup=markup)


async def envoyer_rappel(bot, rappel):
    ids = CIBLES[rappel["cible"]][1]()
    markup = clavier_rappel(rappel["bouton"])
    ok = erreurs = 0
    for uid in ids:
        try:
            await envoyer_a(bot, uid, rappel["texte"], rappel["video"], markup)
            ok += 1
        except Exception as e:
            erreurs += 1
            print(f"Rappel #{rappel['id']} uid={uid} : {e}")
        await asyncio.sleep(PAUSE_ENVOI)
    await prevenir_admins(bot, f"✅ Rappel #{rappel['id']} « {rappel['nom']} » : "
                               f"{ok}/{len(ids)} envoyés ({erreurs} échecs : bot bloqué, etc.)")


async def envoyer_relance(bot, u, colonne):
    """+10 min : texte. +30 min : même texte avec la vidéo de bienvenue."""
    n = restantes(u)
    reste = ("il ne te reste qu'un clic pour confirmer ta place" if n == 0
             else f"il te reste {n} question{'s' if n > 1 else ''} pour confirmer ta place")
    texte = f"{u.get('prenom') or 'Hello'}, {reste}.\n\nOù est-ce que tu bloques ?"
    markup = kb(("Je continue", "demarrer"), ("Je suis bloqué(e)", "souci"))
    if colonne == "relance30":
        await envoyer_video(bot, u["telegram_id"], VIDEO_RELANCE, texte, markup)
    else:
        await bot.send_message(chat_id=u["telegram_id"], text=texte, reply_markup=markup)


async def relancer_incomplets(bot, maintenant):
    for colonne, minutes in (("relance10", 10), ("relance30", 30)):
        for u in db.users_a_relancer(WEBINAIRE, colonne, maintenant - timedelta(minutes=minutes)):
            db.marquer_relance(u["telegram_id"], colonne)          # d'abord : jamais deux envois
            try:
                await envoyer_relance(bot, u, colonne)
            except Exception as e:
                print(f"Relance uid={u['telegram_id']} : {e}")
            await asyncio.sleep(PAUSE_ENVOI)


async def tick(bot, maintenant=None):
    """Appelée toutes les 30 s : envoie les rappels dont l'heure est arrivée, puis les relances."""
    maintenant = maintenant or db.now_benin()
    for r in db.rappels_a_traiter(maintenant.strftime(FORMAT)):
        retard = maintenant - datetime.strptime(r["date_envoi"], FORMAT)
        if retard > TOLERANCE:
            db.set_statut_rappel(r["id"], db.MANQUE)
            await prevenir_admins(bot, f"⚠️ Rappel #{r['id']} « {r['nom']} » NON envoyé (heure dépassée de "
                                       f"{int(retard.total_seconds() // 60)} min). Pour l'envoyer quand même : "
                                       f"/envoyer_rappel {r['id']} oui")
        elif A_COMPLETER in r["texte"].upper():
            db.set_statut_rappel(r["id"], db.MANQUE)
            await prevenir_admins(bot, f"⚠️ Rappel #{r['id']} « {r['nom']} » NON envoyé : texte pas complété. "
                                       f"Écris-le avec /modifier {r['id']} ton texte, puis /envoyer_rappel {r['id']} oui")
        else:
            db.set_statut_rappel(r["id"], db.ENVOYE)
            await envoyer_rappel(bot, r)
    await relancer_incomplets(bot, maintenant)


async def boucle_planificateur(bot):
    while True:
        try:
            await tick(bot)
        except Exception as e:
            print(f"Planificateur : {e}")
        await asyncio.sleep(30)


async def demarrer_planificateur(app: Application):
    lancer(boucle_planificateur(app.bot))


# ════════════════════════════════════════════════════════════════════════════
# COMMANDES ADMIN — rappels et codes
# ════════════════════════════════════════════════════════════════════════════

@admin_seulement
async def cmd_aide(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Commandes admin :\n\n"
        "/rappels — liste des messages planifiés\n"
        "/modifier <n°> <texte> — changer le texte d'un rappel ({prenom} = prénom)\n"
        "/tester_rappel <n°> — recevoir le rappel chez toi pour vérifier\n"
        "/envoyer_rappel <n°> oui — l'envoyer maintenant à sa cible\n"
        "/codes 1 NASDAQ OR BTC — définir les 3 codes du soir 1 (ou 2)\n"
        "/envoyer — diffuser un message à une cible\n"
        "/export — exporter un tableau Excel (tu choisis la cible)\n"
        "/stats — statistiques\n"
        "/nouvelle_categorie <nom>")


@admin_seulement
async def cmd_rappels(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    icones = {db.PLANIFIE: "⏳", db.ENVOYE: "✅", db.MANQUE: "⚠️"}
    lignes = []
    for r in db.get_rappels():
        date = datetime.strptime(r["date_envoi"], FORMAT).strftime("%d/%m %H:%M")
        extrait = r["texte"].replace("\n", " ")[:55]
        lignes.append(f"{icones[r['statut']]} #{r['id']} · {date} · {r['cible']}\n      {extrait}")
    await update.message.reply_text("⏳ planifié · ✅ envoyé · ⚠️ non envoyé\n\n" + "\n".join(lignes))


@admin_seulement
async def cmd_modifier(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    parties = update.message.text.split(maxsplit=2)
    if len(parties) < 3 or not parties[1].isdigit() or not db.get_rappel(int(parties[1])):
        await update.message.reply_text("Usage : /modifier <n°> <nouveau texte>\nEx : /modifier 3 Ce soir 21h !")
        return
    db.modifier_texte_rappel(int(parties[1]), parties[2])
    await update.message.reply_text(f"✅ Rappel #{parties[1]} mis à jour. Vérifie avec /tester_rappel {parties[1]}")


@admin_seulement
async def cmd_tester_rappel(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    rappel = ctx.args and ctx.args[0].isdigit() and db.get_rappel(int(ctx.args[0]))
    if not rappel:
        await update.message.reply_text("Usage : /tester_rappel <n°>")
        return
    await update.message.reply_text(f"Aperçu du rappel #{rappel['id']} (cible : {rappel['cible']}) :")
    await envoyer_a(ctx.bot, update.effective_user.id, rappel["texte"], rappel["video"],
                    clavier_rappel(rappel["bouton"]))


@admin_seulement
async def cmd_envoyer_rappel(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    rappel = ctx.args and ctx.args[0].isdigit() and db.get_rappel(int(ctx.args[0]))
    if not rappel or ctx.args[1:] != ["oui"]:
        await update.message.reply_text("Usage : /envoyer_rappel <n°> oui")
        return
    db.set_statut_rappel(rappel["id"], db.ENVOYE)
    await update.message.reply_text(f"📤 Envoi du rappel #{rappel['id']} lancé…")
    lancer(envoyer_rappel(ctx.bot, rappel))


@admin_seulement
async def cmd_codes(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if len(ctx.args) == 4 and ctx.args[0] in ("1", "2"):
        mots = [normaliser(m) for m in ctx.args[1:]]
        db.definir_codes(int(ctx.args[0]), mots)
        await update.message.reply_text(f"✅ Codes du soir {ctx.args[0]} : " + " · ".join(mots))
        return
    actuels = "\n".join(f"Soir {c['jour']} — code {c['rang']} : {c['mot']}" for c in db.get_codes()) or "(aucun)"
    await update.message.reply_text(f"Codes actuels :\n{actuels}\n\nPour définir : /codes 1 NASDAQ OR BTC")


# ════════════════════════════════════════════════════════════════════════════
# COMMANDES ADMIN — /envoyer et /export (même choix de cible)
# ════════════════════════════════════════════════════════════════════════════

CHOIX_CIBLE, RECEVOIR_MESSAGE = range(2)


async def afficher_cibles(update, ctx, action, titre):
    ctx.user_data["action"] = action
    lignes = [titre, ""]
    for numero, (libelle, get_ids) in enumerate(CIBLES.values(), start=1):
        lignes.append(f"{numero} — {libelle} ({len(get_ids())})")
    lignes.append("\nRéponds avec le numéro (ou /cancel).")
    await update.message.reply_text("\n".join(lignes))
    return CHOIX_CIBLE


@admin_seulement
async def cmd_envoyer(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    return await afficher_cibles(update, ctx, "envoyer", "À qui envoyer le message ?")


@admin_seulement
async def cmd_export(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    return await afficher_cibles(update, ctx, "export", "Quelle liste exporter ?")


async def choisir_cible(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    cles = list(CIBLES)
    choix = update.message.text.strip()
    if not choix.isdigit() or not 1 <= int(choix) <= len(cles):
        await update.message.reply_text(f"❌ Réponds avec un numéro entre 1 et {len(cles)}.")
        return CHOIX_CIBLE
    cle = cles[int(choix) - 1]
    ctx.user_data["cible"] = cle

    if ctx.user_data["action"] == "export":
        await exporter(update, cle)
        return ConversationHandler.END
    await update.message.reply_text(f"✅ Cible : {CIBLES[cle][0]}\n\nEnvoie maintenant ton message "
                                    "(texte, photo ou vidéo avec légende).")
    return RECEVOIR_MESSAGE


async def recevoir_message(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    cle = ctx.user_data["cible"]
    await update.message.reply_text("📤 Diffusion lancée en arrière-plan…")
    lancer(diffuser(ctx.bot, update.effective_chat.id, update.message.message_id, cle))
    return ConversationHandler.END


async def annuler(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("❌ Annulé.")
    return ConversationHandler.END


async def diffuser(bot, admin_chat, message_id, cle):
    """Copie le message de l'admin (quel que soit son format) à chaque personne de la cible."""
    libelle, get_ids = CIBLES[cle]
    ids = get_ids()
    ok = 0
    for uid in ids:
        try:
            await bot.copy_message(chat_id=uid, from_chat_id=admin_chat, message_id=message_id)
            ok += 1
        except Exception as e:
            print(f"Diffusion uid={uid} : {e}")
        await asyncio.sleep(PAUSE_ENVOI)
    await bot.send_message(chat_id=admin_chat, text=f"Diffusion terminée — {ok}/{len(ids)} envoyés à « {libelle} ».")


def construire_excel(lignes) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    entetes = list(lignes[0])
    ws.append(entetes)
    for ligne in lignes:
        ws.append([ligne[e] for e in entetes])

    cles = {"Q3 Déjà tradé", "Q4 Frein", "Q5 Présence"}          # colonnes à lire avant le live
    for cellule in ws[1]:
        couleur = "C00000" if cellule.value in cles else "1F4E79"
        cellule.fill = PatternFill("solid", fgColor=couleur)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[cellule.column_letter].width = 22
    ws.freeze_panes = "A2"

    synthese = wb.create_sheet("Synthèse")
    for titre in ("Q3 Déjà tradé", "Q4 Frein", "Q5 Présence"):
        synthese.append([titre])
        synthese.append(["Réponse", "Nombre"])
        for reponse, n in Counter(l[titre] or "(vide)" for l in lignes).most_common():
            synthese.append([reponse, n])
        synthese.append([])
    synthese.column_dimensions["A"].width = 45

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def exporter(update, cle):
    libelle, get_ids = CIBLES[cle]
    ids = get_ids()
    if not ids:
        await update.message.reply_text(f"❌ Personne dans « {libelle} ».")
        return
    lignes = db.lignes_export(ids, SEUILS)
    date = db.now_benin().strftime("%d%m_%Hh%M")
    await update.message.reply_document(
        document=construire_excel(lignes), filename=f"export_{cle}_{date}.xlsx",
        caption=f"Export « {libelle} » : {len(lignes)} personnes.\n"
                "Onglet « Synthèse » : réponses Q3 / Q4 / Q5 comptées.")


# ════════════════════════════════════════════════════════════════════════════
# COMMANDES ADMIN — /stats et catégories
# ════════════════════════════════════════════════════════════════════════════

@admin_seulement
async def cmd_stats(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    def bloc(dictionnaire):
        return "\n".join(f"  • {k} : {v}" for k, v in dictionnaire.items()) or "  (aucun)"

    tous, confirmes = db.ids_tous(), db.ids_confirmes(WEBINAIRE)
    soirs = "\n".join(
        f"  • {JOURS[j]['nom']} : {len(db.ids_confirmes(WEBINAIRE, j))} prévus · "
        f"{db.nb_clics('live', j)} ont cliqué · {len(db.ids_presents(j, SEUILS[j]))} présents"
        for j in JOURS)
    await update.message.reply_text(
        "Statistiques Trading Pour Tous\n\n"
        f"👥 Personnes connues : {len(tous)}\n"
        f"✅ Confirmées : {len(confirmes)}\n"
        f"⏳ Pas encore confirmées : {len(tous) - len(confirmes)}\n\n"
        f"Liens d'entrée :\n{bloc(db.compter_liens())}\n\n"
        f"Q3 Déjà tradé :\n{bloc(db.compter_reponses(WEBINAIRE, 'deja_trade'))}\n\n"
        f"Q4 Freins :\n{bloc(db.compter_reponses(WEBINAIRE, 'frein'))}\n\n"
        f"Q5 Présence :\n{bloc(db.compter_reponses(WEBINAIRE, 'presence'))}\n\n"
        f"Soirs :\n{soirs}\n\n"
        f"📅 Ont cliqué sur le calendrier : {db.nb_clics('agenda_android') + db.nb_clics('agenda_iphone')}")


@admin_seulement
async def cmd_categorie(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    nom = " ".join(ctx.args).strip()
    if len(nom) < 3:
        liste = "\n".join(f"• {c}" for c in db.get_categories())
        await update.message.reply_text(f"Catégories :\n{liste}\n\nPour en créer une : /nouvelle_categorie <nom>")
        return
    db.ajouter_categorie(nom)
    await update.message.reply_text(f"✅ Catégorie « {nom} » créée.")


# ════════════════════════════════════════════════════════════════════════════
# ERREURS ET LANCEMENT
# ════════════════════════════════════════════════════════════════════════════

async def error_handler(update: object, ctx: ContextTypes.DEFAULT_TYPE):
    tb = "".join(traceback.format_exception(type(ctx.error), ctx.error, ctx.error.__traceback__))
    print(f"[ERROR] {tb}")
    qui = ""
    if isinstance(update, Update) and update.effective_user:
        qui = f"User : {update.effective_user.id}\n\n"
    await prevenir_admins(ctx.bot, f"ERREUR BOT\n\n{qui}{tb[-1800:]}")


def main():
    if not TOKEN:
        raise SystemExit("❌ Définis la variable d'environnement BOT_TOKEN (token régénéré via @BotFather).")

    db.init_db()
    db.seed_rappels(PLANNING)
    init_sondage_db()

    app = (Application.builder().token(TOKEN).read_timeout(30).write_timeout(30)
           .post_init(demarrer_planificateur).build())

    app.add_handler(ChatJoinRequestHandler(approuver_demande))
    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("JeMEnregistre", cmd_inscription))
    for nom, fonction in [("aide", cmd_aide), ("stats", cmd_stats), ("rappels", cmd_rappels),
                          ("modifier", cmd_modifier), ("tester_rappel", cmd_tester_rappel),
                          ("envoyer_rappel", cmd_envoyer_rappel), ("codes", cmd_codes),
                          ("nouvelle_categorie", cmd_categorie)]:
        app.add_handler(CommandHandler(nom, fonction))

    app.add_handler(ConversationHandler(
        entry_points=[CommandHandler("envoyer", cmd_envoyer), CommandHandler("export", cmd_export)],
        states={
            CHOIX_CIBLE: [MessageHandler(filters.TEXT & ~filters.COMMAND, choisir_cible)],
            RECEVOIR_MESSAGE: [MessageHandler(filters.ALL & ~filters.COMMAND, recevoir_message)],
        },
        fallbacks=[CommandHandler("cancel", annuler)],
        per_chat=False, per_user=True, allow_reentry=True,
    ))
    app.add_handler(CallbackQueryHandler(boutons, pattern=r"^(q|prenom|live|agenda|demarrer|souci)(:|$)"))
    app.add_error_handler(error_handler)

    register_sondage_handlers(app)

    # En dernier : capte tout texte hors commande (codes du soir, questionnaire, messages libres)
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, message_libre))

    print("start...")
    app.run_polling(poll_interval=1, allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()